import csv
from pathlib import Path
import unicodedata

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from sistema.models import EnderecoEmpresaMotoboy


NOME_ALIASES = {
    "CLIENTE",
    "EMPRESA",
    "NOME",
    "NOMEDOCLIENTE",
    "RAZAOSOCIAL",
    "RAZAO",
}
ENDERECO_ALIASES = {
    "ENDERECO",
    "ENDERECOCOMPLETO",
    "RUA",
    "LOGRADOURO",
    "LOCAL",
}
NUMERO_ALIASES = {"NUMERO", "N", "NRO"}
BAIRRO_ALIASES = {"BAIRRO"}
CIDADE_ALIASES = {"CIDADE", "MUNICIPIO"}
UF_ALIASES = {"UF", "ESTADO"}


def texto_limpo(valor):
    if valor is None:
        return ""
    return " ".join(str(valor).strip().split())


def texto_chave(valor):
    texto = texto_limpo(valor).upper()
    texto = unicodedata.normalize("NFKD", texto)
    return "".join(ch for ch in texto if not unicodedata.combining(ch) and ch.isalnum())


def coluna_por_alias(colunas, aliases):
    for alias in aliases:
        if alias in colunas:
            return colunas[alias]
    return None


class Command(BaseCommand):
    help = "Importa clientes/endereco para uso nas rotas do motoboy."

    def add_arguments(self, parser):
        parser.add_argument("arquivo", help="Caminho do arquivo .xlsx ou .csv com clientes e enderecos.")
        parser.add_argument(
            "--substituir",
            action="store_true",
            help="Remove todos os clientes de rota antes de importar a planilha.",
        )

    def handle(self, *args, **options):
        caminho = Path(options["arquivo"])
        if not caminho.exists():
            raise CommandError(f"Arquivo nao encontrado: {caminho}")

        linhas = self.ler_linhas(caminho)
        if not linhas:
            raise CommandError("Nenhuma linha encontrada para importar.")

        with transaction.atomic():
            if options["substituir"]:
                EnderecoEmpresaMotoboy.objects.all().delete()

            criados = 0
            atualizados = 0
            ignorados = 0
            for linha in linhas:
                nome = texto_limpo(linha.get("nome"))
                endereco = texto_limpo(linha.get("endereco"))
                if not nome or not endereco:
                    ignorados += 1
                    continue

                cliente, criado = EnderecoEmpresaMotoboy.objects.update_or_create(
                    nome=nome,
                    defaults={"endereco": endereco, "ativo": True},
                )
                if criado:
                    criados += 1
                else:
                    atualizados += 1

        self.stdout.write(
            self.style.SUCCESS(
                f"Clientes importados: {criados} criados, {atualizados} atualizados, {ignorados} ignorados."
            )
        )

    def ler_linhas(self, caminho):
        if caminho.suffix.lower() == ".csv":
            return self.ler_csv(caminho)
        if caminho.suffix.lower() in {".xlsx", ".xlsm"}:
            return self.ler_excel(caminho)
        raise CommandError("Formato nao suportado. Use .xlsx, .xlsm ou .csv.")

    def ler_csv(self, caminho):
        with caminho.open("r", encoding="utf-8-sig", newline="") as arquivo:
            leitor = csv.reader(arquivo, delimiter=";")
            linhas = list(leitor)
        if linhas and len(linhas[0]) == 1:
            with caminho.open("r", encoding="utf-8-sig", newline="") as arquivo:
                linhas = list(csv.reader(arquivo, delimiter=","))
        return self.normalizar_linhas(linhas)

    def ler_excel(self, caminho):
        wb = load_workbook(filename=caminho, data_only=True, read_only=True)
        sheet = wb.active
        linhas = [
            [cell.value for cell in row]
            for row in sheet.iter_rows()
        ]
        return self.normalizar_linhas(linhas)

    def normalizar_linhas(self, linhas):
        linha_cabecalho = None
        colunas = {}
        for indice, linha in enumerate(linhas[:10]):
            mapa = {
                texto_chave(valor): posicao
                for posicao, valor in enumerate(linha)
                if texto_chave(valor)
            }
            if coluna_por_alias(mapa, NOME_ALIASES) is not None and (
                coluna_por_alias(mapa, ENDERECO_ALIASES) is not None
                or coluna_por_alias(mapa, {"LOGRADOURO"}) is not None
            ):
                linha_cabecalho = indice
                colunas = mapa
                break

        if linha_cabecalho is None:
            raise CommandError("Nao encontrei cabecalhos de cliente e endereco na planilha.")

        col_nome = coluna_por_alias(colunas, NOME_ALIASES)
        col_endereco = coluna_por_alias(colunas, ENDERECO_ALIASES)
        col_numero = coluna_por_alias(colunas, NUMERO_ALIASES)
        col_bairro = coluna_por_alias(colunas, BAIRRO_ALIASES)
        col_cidade = coluna_por_alias(colunas, CIDADE_ALIASES)
        col_uf = coluna_por_alias(colunas, UF_ALIASES)

        normalizadas = []
        for linha in linhas[linha_cabecalho + 1:]:
            nome = self.valor_coluna(linha, col_nome)
            endereco = self.valor_coluna(linha, col_endereco)
            partes_endereco = [endereco]
            numero = self.valor_coluna(linha, col_numero)
            if numero and numero not in endereco:
                partes_endereco.append(numero)
            partes_endereco.extend(
                [
                    self.valor_coluna(linha, col_bairro),
                    self.valor_coluna(linha, col_cidade),
                    self.valor_coluna(linha, col_uf),
                ]
            )
            normalizadas.append(
                {
                    "nome": nome,
                    "endereco": ", ".join(parte for parte in partes_endereco if parte),
                }
            )
        return normalizadas

    def valor_coluna(self, linha, coluna):
        if coluna is None or coluna >= len(linha):
            return ""
        return texto_limpo(linha[coluna])
