from datetime import date, datetime, time
from pathlib import Path
from tempfile import TemporaryDirectory
import unicodedata

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.core.management.base import CommandError
from django.db.models import Q
from django.views.decorators.http import require_POST
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from openpyxl import load_workbook

from sistema.models import ChamadoTI, ConfiguracaoChamadosTI, ModuloSistema
from sistema.permissions import usuario_eh_admin, usuario_pode_editar, usuario_tem_acesso
from sistema.utils import enviar_email_abertura_chamado_ti

from ..common import contexto_modulo, renderizar_modulo_sem_permissao
from .forms import (
    ChamadoTIAberturaForm,
    ChamadoTIConclusaoForm,
    ConfiguracaoChamadosTIForm,
    ImportarChamadosTIForm,
)


def _texto(valor):
    if valor is None:
        return ""
    return str(valor).strip()


def _chave(valor):
    texto = unicodedata.normalize("NFKD", _texto(valor).casefold())
    texto = "".join(caractere for caractere in texto if not unicodedata.combining(caractere))
    return "".join(caractere for caractere in texto if caractere.isalnum())


def _normalizar_prioridade(valor):
    chave = _chave(valor)
    if "urgente" in chave:
        return ChamadoTI.Prioridade.URGENTE
    if "alta" in chave:
        return ChamadoTI.Prioridade.ALTA
    if "baixa" in chave:
        return ChamadoTI.Prioridade.BAIXA
    return ChamadoTI.Prioridade.MEDIA


def _normalizar_status(valor):
    chave = _chave(valor)
    if "concluido" in chave or "finalizado" in chave or "resolvido" in chave:
        return ChamadoTI.Status.CONCLUIDO
    if "cancelado" in chave:
        return ChamadoTI.Status.CANCELADO
    if "aguard" in chave:
        return ChamadoTI.Status.AGUARDANDO
    if "andamento" in chave:
        return ChamadoTI.Status.EM_ANDAMENTO
    return ChamadoTI.Status.ABERTO


def _data(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return timezone.localdate()


def _hora(valor):
    if isinstance(valor, datetime):
        return valor.time().replace(second=0, microsecond=0)
    if isinstance(valor, time):
        return valor.replace(second=0, microsecond=0)
    return None


def _combinar_data_hora(data_base, hora_base):
    if not hora_base:
        return None
    valor = datetime.combine(data_base, hora_base)
    if timezone.is_naive(valor):
        return timezone.make_aware(valor, timezone.get_current_timezone())
    return valor


def _inteiro(valor):
    if valor in (None, ""):
        return None
    try:
        return max(0, int(float(valor)))
    except (TypeError, ValueError):
        return None


def _proximo_codigo_chamado():
    ultimo_codigo = (
        ChamadoTI.objects.exclude(codigo="")
        .order_by("-codigo")
        .values_list("codigo", flat=True)
        .first()
    )
    try:
        proximo = int(ultimo_codigo) + 1 if ultimo_codigo else 1
    except ValueError:
        proximo = ChamadoTI.objects.count() + 1
    return f"{proximo:04d}"


def _salvar_upload_temporario(arquivo, diretorio):
    caminho = Path(diretorio) / arquivo.name
    with caminho.open("wb") as destino:
        for chunk in arquivo.chunks():
            destino.write(chunk)
    return caminho


def _mapear_colunas(sheet):
    colunas = {}
    for col_idx in range(1, sheet.max_column + 1):
        chave = _chave(sheet.cell(row=2, column=col_idx).value)
        if chave:
            colunas[chave] = col_idx
    return colunas


def _valor(sheet, row_idx, colunas, nome):
    coluna = _coluna(colunas, nome)
    if not coluna:
        return None
    return sheet.cell(row=row_idx, column=coluna).value


def _coluna(colunas, nome):
    esperado = _chave(nome)
    if esperado in colunas:
        return colunas[esperado]
    partes = [parte for parte in esperado.split() if parte]
    for chave, coluna in colunas.items():
        if chave == esperado or esperado in chave or chave in esperado:
            return coluna
    if "descricao" in esperado or "solicitacao" in esperado:
        for chave, coluna in colunas.items():
            if chave.startswith("descr") or "solicit" in chave:
                return coluna
    if "solucao" in esperado:
        for chave, coluna in colunas.items():
            if chave.startswith("solu") or "aplicada" in chave:
                return coluna
    if "observacoes" in esperado:
        for chave, coluna in colunas.items():
            if chave.startswith("observ"):
                return coluna
    if "horainicio" in esperado:
        for chave, coluna in colunas.items():
            if "hora" in chave and ("inicio" in chave or "incio" in chave):
                return coluna
    if "horafim" in esperado:
        for chave, coluna in colunas.items():
            if "hora" in chave and "fim" in chave:
                return coluna
    if "tempomin" in esperado:
        for chave, coluna in colunas.items():
            if "tempo" in chave and "min" in chave:
                return coluna
    return None


def _importar_planilha_chamados(caminho, substituir, usuario):
    wb = load_workbook(filename=caminho, data_only=True)
    if "Atendimentos" not in wb.sheetnames:
        raise CommandError("A aba Atendimentos nao foi encontrada.")
    sheet = wb["Atendimentos"]
    colunas = _mapear_colunas(sheet)
    obrigatorias = ["ID", "Data", "Colaborador", "Descricao da solicitacao"]
    if any(not _coluna(colunas, nome) for nome in obrigatorias):
        raise CommandError("A planilha nao possui as colunas obrigatorias de chamados.")

    if substituir:
        ChamadoTI.objects.all().delete()

    resumo = {"criados": 0, "atualizados": 0, "ignorados": 0}
    for row_idx in range(3, sheet.max_row + 1):
        codigo = _texto(_valor(sheet, row_idx, colunas, "ID"))
        colaborador = _texto(_valor(sheet, row_idx, colunas, "Colaborador"))
        descricao = _texto(_valor(sheet, row_idx, colunas, "Descricao da solicitacao"))
        if not codigo or not colaborador or not descricao:
            resumo["ignorados"] += 1
            continue

        chamado, criado = ChamadoTI.objects.get_or_create(
            codigo=codigo,
            defaults={"criado_por": usuario},
        )
        chamado.data = _data(_valor(sheet, row_idx, colunas, "Data"))
        chamado.atendente_nome = _texto(_valor(sheet, row_idx, colunas, "Atendente"))
        chamado.setor = _texto(_valor(sheet, row_idx, colunas, "Setor"))[:80]
        chamado.colaborador = colaborador[:120]
        chamado.prioridade = _normalizar_prioridade(_valor(sheet, row_idx, colunas, "Prioridade"))
        chamado.status = _normalizar_status(_valor(sheet, row_idx, colunas, "Status"))
        chamado.descricao = descricao
        chamado.solucao = _texto(_valor(sheet, row_idx, colunas, "Solucao aplicada"))
        chamado.hora_inicio = _hora(_valor(sheet, row_idx, colunas, "Hora inicio"))
        chamado.hora_fim = _hora(_valor(sheet, row_idx, colunas, "Hora fim"))
        chamado.tempo_minutos = _inteiro(_valor(sheet, row_idx, colunas, "Tempo min"))
        chamado.aberto_em = _combinar_data_hora(chamado.data, chamado.hora_inicio) or timezone.now()
        chamado.atendimento_iniciado_em = _combinar_data_hora(chamado.data, chamado.hora_inicio)
        chamado.concluido_em = _combinar_data_hora(chamado.data, chamado.hora_fim)
        chamado.tempo_atendimento_minutos = chamado.tempo_minutos
        chamado.observacoes = _texto(_valor(sheet, row_idx, colunas, "Observacoes"))
        chamado.save()
        resumo["criados" if criado else "atualizados"] += 1
    return resumo


@login_required
def chamados_ti(request):
    modulo = ModuloSistema.CHAMADOS_TI
    if not usuario_tem_acesso(request.user, modulo):
        return renderizar_modulo_sem_permissao(
            request,
            modulo,
            "Chamados - TI",
            "Fila de atendimento, prioridades e acompanhamento de resolucao.",
            "sistema/chamados_ti.html",
        )

    busca = request.GET.get("q", "").strip()
    status = request.GET.get("status", "").strip()
    prioridade = request.GET.get("prioridade", "").strip()
    chamados = ChamadoTI.objects.all()
    if busca:
        chamados = chamados.filter(
            Q(codigo__icontains=busca)
            | Q(colaborador__icontains=busca)
            | Q(setor__icontains=busca)
            | Q(descricao__icontains=busca)
            | Q(solucao__icontains=busca)
        )
    if status in ChamadoTI.Status.values:
        chamados = chamados.filter(status=status)
    else:
        status = ""
    if prioridade in ChamadoTI.Prioridade.values:
        chamados = chamados.filter(prioridade=prioridade)
    else:
        prioridade = ""

    context = contexto_modulo(
        request,
        modulo,
        "Chamados - TI",
        "Registre, priorize e acompanhe solicitacoes recebidas por WhatsApp, e-mail ou atendimento interno.",
        extra={
            "chamados": chamados[:120],
            "filtros": {"q": busca, "status": status, "prioridade": prioridade},
            "status_choices": ChamadoTI.Status.choices,
            "prioridade_choices": ChamadoTI.Prioridade.choices,
            "permite_configurar_emails": usuario_eh_admin(request.user, modulo),
            "total_chamados": ChamadoTI.objects.count(),
            "total_abertos": ChamadoTI.objects.exclude(
                status__in=[ChamadoTI.Status.CONCLUIDO, ChamadoTI.Status.CANCELADO]
            ).count(),
            "total_urgentes": ChamadoTI.objects.filter(prioridade=ChamadoTI.Prioridade.URGENTE).exclude(
                status__in=[ChamadoTI.Status.CONCLUIDO, ChamadoTI.Status.CANCELADO]
            ).count(),
            "total_concluidos": ChamadoTI.objects.filter(status=ChamadoTI.Status.CONCLUIDO).count(),
        },
    )
    return render(request, "sistema/chamados_ti.html", context)


@login_required
def novo_chamado_ti(request):
    modulo = ModuloSistema.CHAMADOS_TI
    if not usuario_tem_acesso(request.user, modulo):
        return renderizar_modulo_sem_permissao(
            request,
            modulo,
            "Novo chamado",
            "Cadastro de chamados de TI.",
            "sistema/chamado_ti_form.html",
            extra={"form": ChamadoTIAberturaForm()},
        )

    if request.method == "POST":
        form = ChamadoTIAberturaForm(request.POST, request.FILES)
        if form.is_valid():
            chamado = form.save(commit=False)
            chamado.codigo = _proximo_codigo_chamado()
            chamado.criado_por = request.user
            if not chamado.colaborador:
                chamado.colaborador = request.user.get_full_name() or request.user.email or request.user.username
            chamado.status = ChamadoTI.Status.ABERTO
            chamado.data = timezone.localdate()
            chamado.aberto_em = timezone.now()
            chamado.save()
            try:
                destinatarios = ConfiguracaoChamadosTI.carregar().listar_emails_abertura()
                enviar_email_abertura_chamado_ti(chamado, destinatarios)
            except Exception as erro:
                messages.warning(
                    request,
                    f"Chamado {chamado.codigo} cadastrado, mas houve falha no envio do e-mail: {erro}.",
                )
            else:
                messages.success(request, f"Chamado {chamado.codigo} cadastrado com sucesso e e-mail enviado.")
            return redirect("chamados_ti")
    else:
        solicitante = request.user.get_full_name() or request.user.email or request.user.username
        form = ChamadoTIAberturaForm(initial={"colaborador": solicitante})

    context = contexto_modulo(
        request,
        modulo,
        "Novo chamado",
        "Registre uma solicitação recebida por WhatsApp, e-mail ou conversa interna.",
        extra={"form": form},
    )
    return render(request, "sistema/chamado_ti_form.html", context)


@login_required
def editar_chamado_ti(request, pk):
    modulo = ModuloSistema.CHAMADOS_TI
    chamado = get_object_or_404(ChamadoTI, pk=pk)
    if not usuario_tem_acesso(request.user, modulo):
        return renderizar_modulo_sem_permissao(
            request,
            modulo,
            "Editar chamado",
            "Atualizacao de chamado de TI.",
            "sistema/chamado_ti_form.html",
            extra={"form": ChamadoTIAberturaForm(instance=chamado), "chamado": chamado},
        )
    if not usuario_pode_editar(request.user, modulo):
        messages.error(request, "Seu perfil permite apenas visualizar os chamados.")
        return redirect("chamados_ti")

    if request.method == "POST":
        form = ChamadoTIAberturaForm(request.POST, request.FILES, instance=chamado)
        if form.is_valid():
            chamado = form.save()
            messages.success(request, f"Chamado {chamado.codigo} atualizado com sucesso.")
            return redirect("chamados_ti")
    else:
        form = ChamadoTIAberturaForm(instance=chamado)

    context = contexto_modulo(
        request,
        modulo,
        f"Chamado {chamado.codigo}",
        "Atualize os dados da solicitação e observações internas.",
        extra={"form": form, "chamado": chamado},
    )
    return render(request, "sistema/chamado_ti_form.html", context)


@login_required
@require_POST
def assumir_chamado_ti(request, pk):
    modulo = ModuloSistema.CHAMADOS_TI
    chamado = get_object_or_404(ChamadoTI, pk=pk)
    if not usuario_pode_editar(request.user, modulo):
        messages.error(request, "Seu perfil permite apenas visualizar os chamados.")
        return redirect("chamados_ti")
    if chamado.status == ChamadoTI.Status.CONCLUIDO:
        messages.warning(request, f"Chamado {chamado.codigo} ja esta concluido.")
        return redirect("chamados_ti")
    chamado.registrar_inicio_atendimento(request.user)
    chamado.save(update_fields=["status", "atendente", "atendente_nome", "atendimento_iniciado_em", "atualizado_em"])
    messages.success(request, f"Chamado {chamado.codigo} assumido para atendimento.")
    return redirect("chamados_ti")


@login_required
def concluir_chamado_ti(request, pk):
    modulo = ModuloSistema.CHAMADOS_TI
    chamado = get_object_or_404(ChamadoTI, pk=pk)
    if not usuario_tem_acesso(request.user, modulo):
        return renderizar_modulo_sem_permissao(
            request,
            modulo,
            "Concluir chamado",
            "Finalização de chamado de TI.",
            "sistema/chamado_ti_concluir.html",
            extra={"form": ChamadoTIConclusaoForm(instance=chamado), "chamado": chamado},
        )
    if not usuario_pode_editar(request.user, modulo):
        messages.error(request, "Seu perfil permite apenas visualizar os chamados.")
        return redirect("chamados_ti")

    if request.method == "POST":
        form = ChamadoTIConclusaoForm(request.POST, instance=chamado)
        if form.is_valid():
            chamado.registrar_conclusao(form.cleaned_data["solucao"], request.user)
            chamado.save(
                update_fields=[
                    "status",
                    "atendente",
                    "atendente_nome",
                    "solucao",
                    "atendimento_iniciado_em",
                    "concluido_em",
                    "tempo_atendimento_minutos",
                    "atualizado_em",
                ]
            )
            messages.success(request, f"Chamado {chamado.codigo} concluído com sucesso.")
            return redirect("chamados_ti")
    else:
        form = ChamadoTIConclusaoForm(instance=chamado)

    context = contexto_modulo(
        request,
        modulo,
        f"Concluir chamado {chamado.codigo}",
        "Informe a solução aplicada para finalizar o atendimento.",
        extra={"form": form, "chamado": chamado},
    )
    return render(request, "sistema/chamado_ti_concluir.html", context)


@login_required
def configurar_emails_chamados_ti(request):
    modulo = ModuloSistema.CHAMADOS_TI
    if not usuario_eh_admin(request.user, modulo):
        messages.error(request, "Apenas administradores de Chamados - TI podem alterar estes e-mails.")
        return redirect("chamados_ti")

    configuracao = ConfiguracaoChamadosTI.carregar()
    if request.method == "POST":
        form = ConfiguracaoChamadosTIForm(request.POST, instance=configuracao)
        if form.is_valid():
            configuracao = form.save(commit=False)
            configuracao.atualizado_por = request.user
            configuracao.save()
            messages.success(request, "Destinatários de abertura atualizados com sucesso.")
            return redirect("chamados_ti")
    else:
        form = ConfiguracaoChamadosTIForm(instance=configuracao)

    context = contexto_modulo(
        request,
        modulo,
        "E-mails de abertura",
        "Defina quem recebe o aviso quando um chamado de TI for aberto.",
        extra={"form": form, "configuracao": configuracao},
    )
    return render(request, "sistema/chamados_ti_configuracao.html", context)


@login_required
def importar_chamados_ti(request):
    modulo = ModuloSistema.CHAMADOS_TI
    context = contexto_modulo(
        request,
        modulo,
        "Importar chamados",
        "Migre a planilha atual de atendimentos para a fila de chamados do sistema.",
    )
    if context["acesso_negado"]:
        context["form"] = ImportarChamadosTIForm()
        return render(request, "sistema/chamados_ti_importar.html", context)
    if not context["permite_edicao"]:
        messages.error(request, "Seu perfil permite apenas visualizar os chamados.")
        return redirect("chamados_ti")

    if request.method == "POST":
        form = ImportarChamadosTIForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                with TemporaryDirectory() as diretorio:
                    caminho = _salvar_upload_temporario(form.cleaned_data["arquivo"], diretorio)
                    resumo = _importar_planilha_chamados(
                        caminho,
                        form.cleaned_data["substituir_chamados"],
                        request.user,
                    )
            except (CommandError, ValidationError) as erro:
                messages.error(request, f"Nao foi possivel importar a planilha: {erro}")
            else:
                messages.success(
                    request,
                    "Chamados importados com sucesso: "
                    f"{resumo['criados']} criados, {resumo['atualizados']} atualizados, "
                    f"{resumo['ignorados']} ignorados.",
                )
                return redirect("chamados_ti")
    else:
        form = ImportarChamadosTIForm()

    context["form"] = form
    return render(request, "sistema/chamados_ti_importar.html", context)
